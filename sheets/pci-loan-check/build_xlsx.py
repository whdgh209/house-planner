# -*- coding: utf-8 -*-
"""PCI_차용검증.xlsx 생성 — Google Sheets 가져오기용.

IFS/TEXTJOIN 등 post-2007 함수는 쓰지 않는다. 중첩 IF와 문자열 연결만 사용해
Excel / LibreOffice / Google Sheets 모두에서 동일하게 동작하게 한다.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule, ColorScaleRule
from openpyxl.utils import get_column_letter

WON, PCT1, PCT2 = '#,##0"원"', '0.0%', '0.00%'
MONTH, TIMES, DATE, YEAR = '#,##0"개월"', '0.00"배"', 'yyyy-mm-dd', '0"년"'
PLAIN = '#,##0'

F_TITLE = PatternFill('solid', fgColor='1F3864')
F_HEAD  = PatternFill('solid', fgColor='D9E2F3')
F_IN    = PatternFill('solid', fgColor='FFF2CC')
F_CALC  = PatternFill('solid', fgColor='F2F2F2')
F_OUT   = PatternFill('solid', fgColor='E2EFDA')
F_RISK  = PatternFill('solid', fgColor='FCE4E4')

ARIAL = 'Arial'
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
ws = wb.active
ws.title = 'PCI_차용검증'

def put(r, label, value=None, note=None, style=None, fmt=None):
    ws.cell(r, 1, label)
    if value is not None:
        c = ws.cell(r, 2, value)
        if fmt:
            c.number_format = fmt
    if note is not None:
        n = ws.cell(r, 3, note)
        n.font = Font(name=ARIAL, size=9, color='7F7F7F')
    if style == 'head':
        for col in range(1, 4):
            ws.cell(r, col).fill = F_HEAD
            ws.cell(r, col).font = Font(name=ARIAL, bold=True)
    elif style == 'input':
        ws.cell(r, 2).fill = F_IN
        ws.cell(r, 2).font = Font(name=ARIAL, bold=True)
    elif style == 'calc':
        ws.cell(r, 2).fill = F_CALC
    elif style == 'out':
        ws.cell(r, 2).fill = F_OUT
        ws.cell(r, 2).font = Font(name=ARIAL, bold=True)
    elif style == 'risk':
        ws.cell(r, 2).fill = F_RISK
        ws.cell(r, 2).font = Font(name=ARIAL, bold=True)

# ── 제목 ────────────────────────────────────────────────────────────────
ws.merge_cells('A1:I1')
t = ws['A1']
t.value = '개인 간 차용 및 상환능력 검증 (PCI 세무 리스크 사전 점검)'
t.fill = F_TITLE
t.font = Font(name=ARIAL, size=13, bold=True, color='FFFFFF')
t.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 34

put(2, '작성 기준일', '=TODAY()', None, 'calc', DATE)
put(3, '차입자 / 대여자', '봄 / 종호', '차입자 = 돈을 빌리는 사람', 'input')

# ── 범례 ────────────────────────────────────────────────────────────────
ws['E2'] = '■ 범례 — 읽는 법'
ws['E2'].font = Font(name=ARIAL, bold=True)
legend = [
    ('E3', F_IN,   '입력 셀 — 여기만 고치면 됩니다'),
    ('E4', F_CALC, '자동 계산 — 건드리지 마세요'),
    ('E5', F_OUT,  '판정 결과'),
    ('E6', F_RISK, '세무 리스크 진단'),
]
for addr, fill, text in legend:
    ws[addr].fill = fill
    ws[addr].border = BOX
    col = ws[addr].column + 1
    ws.cell(ws[addr].row, col, text).font = Font(name=ARIAL, size=9)
ws['E8'] = '요율은 B14:B20, 법정 기준값은 B47:B49에서 수정하세요.'
ws['E8'].font = Font(name=ARIAL, size=9, color='7F7F7F')
ws['E9'] = '현재 채워진 숫자는 형식을 보여주는 예시값입니다.'
ws['E9'].font = Font(name=ARIAL, size=9, color='7F7F7F')

# ── 1. 소득 ─────────────────────────────────────────────────────────────
put(5,  '■ 1. 소득 정보  [입력]', None, None, 'head')
put(6,  '기본 연봉 (세전)', 60000000, '근로계약서 기준 — 예시값', 'input', WON)
put(7,  '연간 예상 성과상여금 (세전)', 10000000, '변동 재원 — 판정에서 별도 취급', 'input', WON)
put(8,  '비과세 식대 (월)', 200000, '월 20만원 한도', 'input', WON)
put(9,  '연간 급여 상승률', 0.03, '연도별 시뮬레이션에 반영', 'input', PCT1)
put(10, '성과급 반영률 (시나리오)', 1, '보수적 점검은 50% 또는 0%', 'input', PCT1)
put(11, '식대가 B6에 포함되어 있는가?', '포함', '드롭다운: 포함 / 미포함', 'input')

# ── 2. 세·4대보험 요율 ──────────────────────────────────────────────────
put(13, '■ 2. 세·4대보험 요율  [입력 · 매년 고시값 확인]', None, None, 'head')
put(14, '국민연금 요율 (근로자)', 0.045, None, 'input', PCT2)
put(15, '국민연금 기준소득월액 상한', 6370000, '매년 7월 변동 — 반드시 확인', 'input', WON)
put(16, '건강보험 요율 (근로자)', 0.03545, None, 'input', PCT2)
put(17, '장기요양보험 요율 (건보료 대비)', 0.1295, '건강보험료 × 이 요율', 'input', PCT2)
put(18, '고용보험 요율 (근로자)', 0.009, None, 'input', PCT2)
put(19, '기본급 근로소득세 실효세율 (지방세 포함)', 0.06, '원천징수영수증 결정세액 ÷ 총급여로 역산 권장', 'input', PCT1)
put(20, '상여금 적용 세율 (지방세 포함)', 0.385, '과표 35% 구간 + 지방소득세 3.5%', 'input', PCT1)

# ── 3. 금융 부채 ────────────────────────────────────────────────────────
put(22, '■ 3. 금융 부채 상환 정보  [입력]', None, None, 'head')
put(23, '주택담보대출 월 상환액', 1200000, '원리금균등 기준', 'input', WON)
put(24, '  └ 잔여 상환 개월수', 300, '만기까지 남은 개월', 'input', MONTH)
put(25, '기타 대출 A 월 상환액 (회사 대출)', 500000, '원금+이자', 'input', WON)
put(26, '  └ 잔여 상환 개월수', 60, None, 'input', MONTH)
put(27, '기타 대출 B 월 상환액 (신협/신용)', 300000, '원금+이자', 'input', WON)
put(28, '  └ 잔여 상환 개월수', 36, None, 'input', MONTH)
put(29, '월 부채상환 합계 (현재)', '=B23+B25+B27', None, 'calc', WON)

# ── 4. 생활비 ───────────────────────────────────────────────────────────
put(31, '■ 4. 차입자 필수 생활비  [입력]', None, None, 'head')
put(32, '월 고정비 (통신·보험·교통 등)', 600000, '차입자 명의 카드/계좌 기준', 'input', WON)
put(33, '월 변동비 (식대·카페·개인 소비)', 900000, '최근 3개월 카드 명세 평균 권장', 'input', WON)
put(34, '월 최소 개인 소비액', '=B32+B33', None, 'calc', WON)
put(35, '연간 소비 증가율', 0.02, '물가 반영', 'input', PCT1)

# ── 5. 차용 조건 ────────────────────────────────────────────────────────
put(37, '■ 5. 개인 간 차용 조건  [입력]', None, None, 'head')
put(38, '차용 원금', 200000000, None, 'input', WON)
put(39, '차용 개시일', '2026-01-01', '실제 계좌이체일', 'input', DATE)
put(40, '차용 기간 (년)', 5, '1~10년 (10년 초과 시 표 행 추가 필요)', 'input', YEAR)
put(41, '약정 이자율 (연)', 0, '무이자면 0%', 'input', PCT1)
put(42, '상환 방식', '만기 일시 상환', '드롭다운: 만기 일시 상환 / 원금 분할 상환', 'input')
put(43, '연간 원금 분할상환액', 0, '만기 일시면 0', 'input', WON)
put(44, '만기일', '=IF(ISNUMBER(B39),EDATE(B39,B40*12),"")', None, 'calc', DATE)

# ── 6. 세법 기준값 ──────────────────────────────────────────────────────
put(46, '■ 6. 세법 기준값  [입력 · 상증세법]', None, None, 'head')
put(47, '법정 적정 이자율', 0.046, '상증세법 시행령 §31조의4', 'input', PCT1)
put(48, '증여이익 과세 기준금액 (연)', 10000000, '이 금액 이상이면 차액 전액 과세', 'input', WON)
put(49, '이자소득 원천징수율 (비영업대금의 이익)', 0.275, '소득세 25% + 지방소득세 2.5%', 'input', PCT1)

# ── 7. 중간 계산 ────────────────────────────────────────────────────────
put(51, '■ 7. 중간 계산  [자동]', None, None, 'head')
put(52, '과세대상 월 급여', '=IF($B$11="포함",($B$6-$B$8*12)/12,$B$6/12)', '식대 비과세분 제외', 'calc', WON)
put(53, '월 국민연금', '=ROUND(MIN($B$52,$B$15)*$B$14,0)', '기준소득월액 상한 적용', 'calc', WON)
put(54, '월 건강보험', '=ROUND($B$52*$B$16,0)', None, 'calc', WON)
put(55, '월 장기요양보험', '=ROUND($B$54*$B$17,0)', None, 'calc', WON)
put(56, '월 고용보험', '=ROUND($B$52*$B$18,0)', None, 'calc', WON)
put(57, '월 4대보험 합계', '=SUM($B$53:$B$56)', None, 'calc', WON)
put(58, '월 근로소득세 + 지방소득세', '=ROUND($B$52*$B$19,0)', None, 'calc', WON)
put(59, '월 세후 실수령액 (식대 포함)', '=$B$52-$B$57-$B$58+$B$8', None, 'calc', WON)
put(60, '연간 기본급 세후 실수령액', '=$B$59*12', None, 'calc', WON)
put(61, '성과급 세후 실수령액 (연)', '=ROUND($B$7*$B$10*(1-$B$20),0)', '시나리오 반영률 적용', 'calc', WON)
put(62, '연간 세후 총소득', '=$B$60+$B$61', None, 'calc', WON)
put(63, '연간 기존 부채 상환액', '=$B$29*12', None, 'calc', WON)
put(64, '연간 차용 이자 (총액, 세전)', '=ROUND($B$38*$B$41,0)', None, 'calc', WON)
put(65, '  └ 원천징수세액 (차입자가 세무서 납부)', '=ROUND($B$64*$B$49,0)', '지급월 익월 10일까지', 'calc', WON)
put(66, '  └ 대여자 실수령 이자', '=$B$64-$B$65', None, 'calc', WON)
put(67, '연간 필수 소비 지출', '=$B$34*12', None, 'calc', WON)
put(68, '연간 순 가처분 잉여현금', '=$B$62-($B$63+$B$64+$B$67)', '1년차 기준', 'calc', WON)
put(69, '월 기본급 기준 수지 (성과급 제외)', '=$B$59-$B$29-$B$34-$B$64/12', '음수면 구조적 적자', 'calc', WON)
put(70, '차용기간 누적 저축 가능액 (단순 추정)', '=$B$68*$B$40', '성장률 미반영 — 참고용', 'calc', WON)

# ── 8. 연도별 시뮬레이션 ────────────────────────────────────────────────
ws['A72'] = '■ 8. 연도별 현금흐름 시뮬레이션  [자동]'
for col in range(1, 10):
    ws.cell(72, col).fill = F_HEAD
    ws.cell(72, col).font = Font(name=ARIAL, bold=True)

hdr = ['연차', '세후 총소득', '기존 부채상환', '필수 생활비', '차용 이자',
       '원금 분할상환', '연간 잉여', '누적 잉여', '원금 대비 누적']
for i, h in enumerate(hdr, start=1):
    c = ws.cell(73, i, h)
    c.fill = F_CALC
    c.font = Font(name=ARIAL, bold=True)
    c.alignment = Alignment(horizontal='center', wrap_text=True)

for n in range(1, 11):
    r = 73 + n
    g = '$A{r}>$B$40'.format(r=r)
    a = ws.cell(r, 1, n)
    a.number_format = YEAR
    a.alignment = Alignment(horizontal='center')
    a.font = Font(name=ARIAL, bold=True)

    ws.cell(r, 2, '=IF({g},"",ROUND($B$62*(1+$B$9)^($A{r}-1),0))'.format(g=g, r=r))
    ws.cell(r, 3, ('=IF({g},"",'
                   '$B$23*MIN(12,MAX(0,$B$24-($A{r}-1)*12))+'
                   '$B$25*MIN(12,MAX(0,$B$26-($A{r}-1)*12))+'
                   '$B$27*MIN(12,MAX(0,$B$28-($A{r}-1)*12)))').format(g=g, r=r))
    ws.cell(r, 4, '=IF({g},"",ROUND($B$67*(1+$B$35)^($A{r}-1),0))'.format(g=g, r=r))
    ws.cell(r, 5, '=IF({g},"",ROUND(MAX(0,$B$38-$B$43*($A{r}-1))*$B$41,0))'.format(g=g, r=r))
    ws.cell(r, 6, '=IF({g},"",MIN($B$43,MAX(0,$B$38-$B$43*($A{r}-1))))'.format(g=g, r=r))
    ws.cell(r, 7, '=IF({g},"",B{r}-C{r}-D{r}-E{r}-F{r})'.format(g=g, r=r))
    if n == 1:
        ws.cell(r, 8, '=IF({g},"",G{r})'.format(g=g, r=r))
    else:
        ws.cell(r, 8, '=IF({g},"",N(H{p})+G{r})'.format(g=g, r=r, p=r - 1))
    ws.cell(r, 9, '=IF({g},"",IFERROR(H{r}/$B$38,""))'.format(g=g, r=r))

    for col in range(2, 8):
        ws.cell(r, col).number_format = PLAIN
    ws.cell(r, 9).number_format = PCT1

for r in range(73, 84):
    for col in range(1, 10):
        ws.cell(r, col).border = BOX

# ── 9. 판정 ─────────────────────────────────────────────────────────────
put(85, '■ 9. 상환능력 판정  [핵심 산출물]', None, None, 'head')
put(86, '만기 시 일시상환 대상 잔여 원금', '=MAX(0,$B$38-SUM($F$74:$F$83))', '분할상환분 차감 후', 'out', WON)
put(87, '차용기간 누적 저축 가능액',
    '=IF($B$40>10,"⚠ 기간 10년 초과 — 시뮬레이션 행 추가 필요",INDEX($H$74:$H$83,$B$40))',
    '연도별 표 기준', 'out', WON)
put(88, '지표1. 만기 상환 버퍼율',
    '=IF(NOT(ISNUMBER($B$87)),"-",IF($B$86=0,"원금 전액 분할상환 완료",$B$87/$B$86))',
    '100% = 딱 맞음, 120%↑ 권장', 'out', PCT1)
put(89, '지표2. 최종 판정',
    '=IF(NOT(ISNUMBER($B$87)),"입력 확인 필요",'
    'IF(AND($B$87>=$B$86*1.2,$B$92>=0),"✅ 상환 여유 충분",'
    'IF(AND($B$87>=$B$86,$B$92>=0),"⚠ 상환 가능 (주의 요망)",'
    '"❌ 상환 능력 부족 (증여 추정 위험)")))',
    '기본급 수지 적자면 무조건 부족 판정', 'out')
put(90, '소득 대비 차입 배수', '=IFERROR($B$38/$B$62,"")', '연 세후소득의 몇 배를 빌리는가', 'calc', TIMES)
put(91, '총부채 상환비율 (생활 DSR)', '=IFERROR(($B$63+$B$64)/$B$62,"")', '세후소득 대비 부채상환 비중', 'calc', PCT1)
put(92, '월 기본급 기준 수지', '=$B$69', '성과급 없이도 흑자여야 안전', 'calc', WON)

# ── 10. 세무 리스크 ─────────────────────────────────────────────────────
put(94, '■ 10. 세무 리스크 진단  [자동]', None, None, 'head')
put(95, '법정 적정이자 (연)', '=ROUND($B$38*$B$47,0)', '원금 × 4.6%', 'calc', WON)
put(96, '실제 약정이자 (연)', '=$B$64', None, 'calc', WON)
put(97, '무상대여 이익 (연)', '=MAX(0,$B$95-$B$96)', '적정이자 − 약정이자', 'risk', WON)
put(98, '증여세 과세대상 여부',
    '=IF($B$97>=$B$48,"❌ 과세 대상 — 이익 전액이 증여재산가액",'
    '"✅ 비과세 (연 "&TEXT($B$48,"#,##0")&"원 미만)")',
    '초과분이 아닌 전액 과세', 'risk')
put(99, '무이자 시 안전 원금 한도', '=ROUND($B$48/$B$47,0)', '1,000만 ÷ 4.6%', 'calc', WON)
put(100, '한도 대비 여유액', '=$B$99-$B$38', '음수면 무이자 구조 불가', 'risk', WON)
put(101, '이자 원천징수 신고 의무',
     '=IF($B$41>0,"있음 — 지급월의 다음 달 10일까지 원천세 신고·납부","없음 (무이자)")',
     '신고 이력이 곧 차용 사실의 증거', 'calc')

comment_parts = [
    ('$B$97>=$B$48',
     '"⚠ 무상대여 이익 "&TEXT($B$97,"#,##0")&"원이 연 "&TEXT($B$48,"#,##0")'
     '&"원 기준 이상 → 초과분이 아닌 이익 전액이 증여재산가액으로 과세될 수 있음."'),
    ('AND($B$41=0,$B$38>$B$99)',
     '"⚠ 무이자 차용 원금이 안전 한도 "&TEXT($B$99,"#,##0")'
     '&"원을 초과. 원금을 낮추거나 최소 이자를 약정할 것."'),
    ('$B$92<0',
     '"⚠ 성과급을 제외한 기본급만으로는 월 "&TEXT(ABS($B$92),"#,##0")'
     '&"원 적자 → 상환 재원을 변동 성과급에 의존. 소명 시 가장 취약한 지점."'),
    ('$B$41>0',
     '"ℹ 이자 지급 시 27.5%를 원천징수해 다음 달 10일까지 납부해야 차용 사실이 인정됨."'),
    ('$B$42="만기 일시 상환"',
     '"ℹ 만기 일시·무이자 구조는 증여 추정이 가장 강한 형태. '
     '매년 일부라도 원금을 이체하면 소명력이 크게 올라감."'),
    ('$B$90>5',
     '"⚠ 차용 원금이 연 세후소득의 "&TEXT($B$90,"0.0")&"배 → 통상 소명 강도가 높아지는 구간."'),
    ('$B$91>0.4',
     '"⚠ 생활 DSR "&TEXT($B$91,"0.0%")&" → 기존 부채만으로 상환 여력이 크게 잠식됨."'),
]
joined = '&'.join('IF({c},{t}&CHAR(10),"")'.format(c=c, t=t) for c, t in comment_parts)
put(102, '지표3. 종합 코멘트', '=' + joined, None, 'risk')
ws['B102'].alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[102].height = 130

# ── 11. 민감도 ──────────────────────────────────────────────────────────
ws['A104'] = '■ 11. 성과급 민감도 분석  [자동]'
for col in range(1, 5):
    ws.cell(104, col).fill = F_HEAD
    ws.cell(104, col).font = Font(name=ARIAL, bold=True)
for i, h in enumerate(['성과급 반영률', '누적 잉여액', '버퍼율', '판정'], start=1):
    c = ws.cell(105, i, h)
    c.fill = F_CALC
    c.font = Font(name=ARIAL, bold=True)
    c.alignment = Alignment(horizontal='center')

for i, s in enumerate([0, 0.5, 1]):
    r = 106 + i
    a = ws.cell(r, 1, s)
    a.number_format = '0%'
    a.alignment = Alignment(horizontal='center')
    a.font = Font(name=ARIAL, bold=True)
    ws.cell(r, 2, '=IF(NOT(ISNUMBER($B$87)),"",ROUND($B$87-($B$10-$A{r})'
                  '*ROUND($B$7*(1-$B$20),0)'
                  '*IF($B$9=0,$B$40,((1+$B$9)^$B$40-1)/$B$9),0))'.format(r=r))
    ws.cell(r, 3, '=IF(OR(NOT(ISNUMBER(B{r})),$B$86=0),"-",B{r}/$B$86)'.format(r=r))
    ws.cell(r, 4, '=IF(NOT(ISNUMBER(B{r})),"",'
                  'IF(B{r}>=$B$86*1.2,"✅ 여유 충분",'
                  'IF(B{r}>=$B$86,"⚠ 가능(주의)","❌ 부족")))'.format(r=r))
    ws.cell(r, 2).number_format = PLAIN
    ws.cell(r, 3).number_format = PCT1
for r in range(105, 109):
    for col in range(1, 5):
        ws.cell(r, col).border = BOX

# ── 12. 체크리스트 ──────────────────────────────────────────────────────
ws['A110'] = '■ 12. 소명 실행 체크리스트  [수동 체크]'
for col in range(1, 4):
    ws.cell(110, col).fill = F_HEAD
    ws.cell(110, col).font = Font(name=ARIAL, bold=True)

items = [
    ('차용증(금전소비대차계약서) 작성', '원금·이자율·만기·상환방법·지연이자 명시'),
    ('차용증에 확정일자 부여', '공증 또는 우체국 내용증명 — 사후 작성 의심 차단'),
    ('대여금 전액을 계좌이체로 수령', '현금 수수는 절대 금지'),
    ('상환은 차입자 본인 계좌 → 대여자 계좌', '제3자 계좌 경유 시 소명 불가'),
    ('이자 지급 시 27.5% 원천징수 후 지급', '무이자면 해당 없음'),
    ('원천징수이행상황신고서 제출', '지급월의 다음 달 10일까지'),
    ('이자소득 지급명세서 제출', '다음 해 2월 말까지'),
    ('매년 일부라도 원금 상환 이력 만들기', '만기 일시 구조의 최대 보완책'),
    ('자금조달계획서에 차입금으로 기재', '주택 취득 시 차용증 첨부'),
    ('대여자의 자금 원천 증빙 확보', '대여자도 자금출처를 소명해야 함'),
]
for i, (label, note) in enumerate(items):
    r = 111 + i
    ws.cell(r, 1, label)
    c = ws.cell(r, 2, '☐')
    c.fill = F_IN
    c.alignment = Alignment(horizontal='center')
    n = ws.cell(r, 3, note)
    n.font = Font(name=ARIAL, size=9, color='7F7F7F')

ws.merge_cells('A122:I122')
d = ws['A122']
d.value = ('※ 본 시트는 사전 점검용 추산입니다. 세율·요율은 고시값 변동에 따라 달라지고, '
           '실제 과세 판단은 개별 사실관계에 좌우되므로 실행 전 세무대리인 확인을 권합니다.')
d.font = Font(name=ARIAL, size=9, color='7F7F7F')
d.alignment = Alignment(wrap_text=True)

# ── 근거 주석 ───────────────────────────────────────────────────────────
notes = {
    'B15': '2025.7~2026.6 적용 국민연금 기준소득월액 상한. 매년 7월 고시 변경 — 국민연금공단 공고 확인 필요.',
    'B16': '2025년 건강보험료율 7.09%의 근로자 부담분 절반.',
    'B17': '2025년 장기요양보험료율 (건강보험료 대비).',
    'B19': '간이세액표 기반 추산 예시값. 정확도를 높이려면 전년도 근로소득 원천징수영수증의 '
           '결정세액 ÷ 총급여로 역산해 입력하세요.',
    'B20': '과세표준 8,800만~1.5억 구간 세율 35% + 지방소득세 3.5% 가정. 본인 과표 구간에 맞게 수정하세요.',
    'B47': '상속세 및 증여세법 시행령 제31조의4 — 금전 무상대출 등에 따른 이익의 증여 시 적정이자율.',
    'B48': '상증세법 §41조의4 — 무상대여 이익이 연 1천만원 이상이면 그 이익 전액이 증여재산가액.',
    'B49': '비영업대금의 이익 원천징수세율. 소득세 25% + 지방소득세 2.5%.',
}
for addr, text in notes.items():
    ws[addr].comment = Comment(text, 'PCI 검증 시트')

# ── 데이터 유효성 ───────────────────────────────────────────────────────
dv1 = DataValidation(type='list', formula1='"포함,미포함"', allow_blank=False)
ws.add_data_validation(dv1); dv1.add(ws['B11'])

dv2 = DataValidation(type='list', formula1='"만기 일시 상환,원금 분할 상환"', allow_blank=False)
ws.add_data_validation(dv2); dv2.add(ws['B42'])

dv3 = DataValidation(type='list', formula1='"☐,☑"', allow_blank=False)
ws.add_data_validation(dv3); dv3.add('B111:B120')

dv4 = DataValidation(type='whole', operator='between', formula1=1, formula2=10,
                     allow_blank=False,
                     error='1~10년 사이로 입력하세요. 10년 초과는 시뮬레이션 행을 추가해야 합니다.')
ws.add_data_validation(dv4); dv4.add(ws['B40'])

dv5 = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0,
                     allow_blank=False, error='0 이상의 숫자를 입력하세요.')
ws.add_data_validation(dv5)
for a in ['B6', 'B7', 'B8', 'B23', 'B25', 'B27', 'B32', 'B33', 'B38', 'B43']:
    dv5.add(ws[a])

# ── 조건부 서식 ─────────────────────────────────────────────────────────
GREEN = PatternFill('solid', start_color='C6EFCE', end_color='C6EFCE')
AMBER = PatternFill('solid', start_color='FFEB9C', end_color='FFEB9C')
RED   = PatternFill('solid', start_color='FFC7CE', end_color='FFC7CE')

for rng in ['B89', 'D106:D108']:
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['ISNUMBER(SEARCH("충분",{0}))'.format(rng.split(':')[0])],
        fill=GREEN, font=Font(color='006100', bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['ISNUMBER(SEARCH("주의",{0}))'.format(rng.split(':')[0])],
        fill=AMBER, font=Font(color='9C6500', bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['ISNUMBER(SEARCH("부족",{0}))'.format(rng.split(':')[0])],
        fill=RED, font=Font(color='9C0006', bold=True)))

for rng in ['B88', 'C106:C108']:
    ws.conditional_formatting.add(rng, CellIsRule(
        operator='greaterThanOrEqual', formula=['1.2'], fill=GREEN))
    ws.conditional_formatting.add(rng, CellIsRule(
        operator='lessThan', formula=['1'], fill=RED))

for rng in ['B68:B70', 'B92', 'B100', 'G74:H83']:
    ws.conditional_formatting.add(rng, CellIsRule(
        operator='lessThan', formula=['0'], font=Font(color='9C0006', bold=True)))

ws.conditional_formatting.add('B98', FormulaRule(
    formula=['ISNUMBER(SEARCH("과세 대상",B98))'],
    fill=RED, font=Font(color='9C0006', bold=True)))

ws.conditional_formatting.add('I74:I83', ColorScaleRule(
    start_type='num', start_value=0,   start_color='F8696B',
    mid_type='num',   mid_value=1,     mid_color='FFEB84',
    end_type='num',   end_value=1.2,   end_color='63BE7B'))

# ── 레이아웃 ────────────────────────────────────────────────────────────
ws.column_dimensions['A'].width = 42
ws.column_dimensions['B'].width = 23
ws.column_dimensions['C'].width = 42
for col in range(4, 10):
    ws.column_dimensions[get_column_letter(col)].width = 17
ws.freeze_panes = 'A2'
ws.sheet_view.showGridLines = False

for row in ws.iter_rows(min_row=1, max_row=122, min_col=1, max_col=9):
    for c in row:
        if c.font is None or c.font.name != ARIAL:
            c.font = Font(name=ARIAL, size=c.font.size or 11, bold=c.font.bold,
                          color=c.font.color)

out = '/home/user/house-planner/sheets/pci-loan-check/PCI_차용검증.xlsx'
wb.save(out)
print('saved:', out)
